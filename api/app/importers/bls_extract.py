"""Create the BLS 4.0 extract CSV from the official XLSX.

The download link of the BLS carries a session token, so the XLSX is fetched by hand
and this script is run once (ADR-0002):

    uv run python -m app.importers.bls_extract <path-to-BLS_4_0_Daten_2025_DE.xlsx>
"""

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.importers.bls import CSV_PATH, FIELDS

CODE_HEADER = "BLS Code"
NAME_HEADER = "Lebensmittelbezeichnung"

# BLS component code of each nutrient column in the extract.
NUTRIENT_CODES = {
    "kcal": "ENERCC",
    "protein": "PROT625",
    "carbs": "CHO",
    "fat": "FAT",
    "fiber": "FIBT",
    "sugar": "SUGAR",
    "sat_fat": "FASAT",
    "salt": "NACL",
}


def _value_column(headers: list[str], code: str) -> int:
    """Find the value column of a component, e.g. "FAT Fett [g/100g]".

    The first word must match exactly (CHO vs. CHOCAL), and the value column is the one
    with a unit; "<code> Datenherkunft" and "<code> Referenz" are skipped.
    """
    matches = [
        index
        for index, header in enumerate(headers)
        if header.split(" ", 1)[0] == code and header.rstrip().endswith("/100g]")
    ]
    if len(matches) != 1:
        raise ValueError(f"expected one value column for {code}, found {len(matches)}")
    return matches[0]


def _format(value: object) -> str:
    """Write numbers compactly; BLS markers such as TR, <LOD or - stay as they are."""
    if value is None:
        return ""
    if not isinstance(value, int | float):
        return str(value).strip()
    return f"{value:.6g}" if value != int(value) else str(int(value))


def extract(xlsx_path: Path, csv_path: Path) -> int:
    """Write the extract and return the number of foods."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    headers = [str(header or "").strip() for header in next(rows)]
    columns = {
        "code": headers.index(CODE_HEADER),
        "name": headers.index(NAME_HEADER),
        **{field: _value_column(headers, code) for field, code in NUTRIENT_CODES.items()},
    }

    count = 0
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            code = row[columns["code"]]
            if not code:
                continue
            record = {"code": str(code).strip(), "name": str(row[columns["name"]]).strip()}
            for field in NUTRIENT_CODES:
                record[field] = _format(row[columns[field]])
            writer.writerow(record)
            count += 1
    workbook.close()
    return count


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("xlsx", type=Path, help="path to BLS_4_0_Daten_2025_DE.xlsx")
    parser.add_argument("--out", type=Path, default=CSV_PATH, help="target CSV")
    args = parser.parse_args(argv)
    count = extract(args.xlsx, args.out)
    print(f"wrote {count} foods to {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
